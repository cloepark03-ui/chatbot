# app.py
import streamlit as st
from openai import OpenAI
from io import BytesIO
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# ------------------ 페이지 설정 ------------------
st.set_page_config(page_title="마케터용 자동 제안서 도우미", page_icon="📊", layout="wide")

# ------------------ 사이드바: 설정 ------------------
with st.sidebar:
    st.header("⚙️ 설정")
    api_key = st.text_input("🔐 OpenAI API Key", type="password", placeholder="sk-...")
    st.caption("키 관리: .streamlit/secrets.toml 사용을 권장")

    st.subheader("🎯 캠페인 기본 정보")
    brand = st.text_input("브랜드명", value="데모브랜드")
    objective = st.multiselect(
        "목표",
        ["인지도", "트래픽", "회원가입/리드", "구매전환", "재구매/CRM"],
        default=["구매전환"]
    )

    industry = st.selectbox(
        "업종",
        ["뷰티/스킨케어", "패션/의류", "F&B/외식", "헬스/의료", "교육/학원", "앱/서비스", "클리닉/병원", "전자/리빙"],
        index=0
    )

    age_range = st.selectbox("핵심 연령대", ["10대", "20대", "30대", "40대", "50대 이상"], index=1)
    gender = st.selectbox("성별", ["전체", "여성", "남성"], index=0)

    st.subheader("🧩 톤 & 매체 프리셋")
    tone = st.selectbox(
        "카피 톤",
        ["데이터 중심/전문적", "감성/스토리텔링", "프리미엄/럭셔리", "경쾌/발랄", "신뢰/권위"],
        index=0
    )
    channels = st.multiselect(
        "매체(프리셋)",
        ["Meta(Instagram/Facebook)", "Naver(검색/GFA)", "Google(Search/Display)", "TikTok", "Kakao Moment"],
        default=["Meta(Instagram/Facebook)", "Naver(검색/GFA)", "Google(Search/Display)"]
    )

    cta_style = st.selectbox("CTA 스타일", ["즉시 행동 유도", "부드러운 제안", "한정/희소성"], index=0)

    st.divider()
    model = st.selectbox("🤖 모델", ["gpt-4o-mini", "gpt-4o", "gpt-3.5-turbo"], index=0)
    temperature = st.slider("창의성(temperature)", 0.0, 1.5, 0.7, 0.1)
    max_tokens = st.slider("Max tokens", 512, 4096, 1800, 64)

# ------------------ 세션 상태 ------------------
if "outputs" not in st.session_state:
    st.session_state.outputs = None

# ------------------ 유틸: 프롬프트 제작 ------------------
def build_prompt(role, brand, industry, age_range, gender, tone, channels, objective, cta_style):
    """role: stp | persona | copy | media"""
    base = f"""
너는 최고 수준의 마케팅 플래너다. 다음 조건에 맞춰 {role.upper()} 결과물을 한국어로 간결하고 구조화해 작성하라.

[브랜드] {brand}
[업종] {industry}
[핵심 타깃] {age_range} / {gender}
[목표] {", ".join(objective)}
[톤] {tone}
[매체] {", ".join(channels)}
[CTA 스타일] {cta_style}
"""
    if role == "stp":
        task = """
STP를 작성하라.
- Segmentation: 3~5개 시장 세그먼트
- Targeting: 우선순위 타깃 1~2개와 선정 근거
- Positioning: 1문장 포지셔닝 + 핵심 근거 3개(USP/RTB)
- 표 형식 요약 표 1개 포함
"""
    elif role == "persona":
        task = """
페르소나를 작성하라(1~2명).
- 이름/나이/직업/라이프스타일
- Pain points(구매장애), Needs, Trigger
- 구매 여정(인지→관심→고려→전환→재방문)에서의 메시지 포인트
- 표 형식 요약 표 1개 포함
"""
    elif role == "copy":
        task = """
카피와 슬로건을 제안하라.
- 인스타그램 광고 주 카피 3개(헤드라인+서브라인+CTA)
- 20대 여성 타깃용 슬로건 3개(업종과 톤 반영)
- 리마케팅용 짧은 헤드라인 5개(12자 내외)
- 해시태그 10개
"""
    else:  # media
        task = """
매체 전략을 작성하라.
- 채널별 역할 정의(퍼널 상단/중단/하단)
- 예산 배분(%) 가이드(초기 4주 기준)
- 캠페인 구조(캠페인/광고세트/소재) 예시
- KPI(CTR, CPC, CPA/ROAS 등)와 측정 이벤트 정의
- 4주 실행 타임라인(주차별 핵심 액션)
- 표 형식 요약 표 1개 포함
"""
    return base + "\n" + task

# ------------------ 생성 함수 ------------------
def generate_all(client, **kwargs):
    system = "너는 상업용 마케팅 제안서 작성 보조 도구다. 간결하고 표/불릿 위주로 출력한다."
    out = {}
    for role in ["stp", "persona", "copy", "media"]:
        prompt = build_prompt(role, **kwargs)
        resp = client.chat.completions.create(
            model=model,
            temperature=temperature,
            max_tokens=max_tokens,
            messages=[
                {"role":"system", "content":system},
                {"role":"user", "content":prompt}
            ],
            stream=False
        )
        out[role] = resp.choices[0].message.content.strip()
    return out

# ------------------ PPT 생성 ------------------
def add_title_slide(prs, title, subtitle):
    slide_layout = prs.slide_layouts[0]  # Title slide
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle
    # 스타일
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(38)
    slide.placeholders[1].text_frame.paragraphs[0].font.size = Pt(18)

def add_text_slide(prs, title, body):
    slide_layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = body
    p.font.size = Pt(16)

def add_two_column_slide(prs, title, left, right):
    slide_layout = prs.slide_layouts[3]  # Two Content
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    left_tf = slide.shapes.placeholders[1].text_frame
    right_tf = slide.shapes.placeholders[2].text_frame
    for tf, txt in [(left_tf, left), (right_tf, right)]:
        tf.clear()
        p = tf.paragraphs[0]
        p.text = txt
        p.font.size = Pt(16)

def build_ppt(brand, industry, objective, outputs):
    prs = Presentation()
    primary = RGBColor(37, 99, 235)

    # 커버
    add_title_slide(
        prs,
        f"{brand} 마케팅 제안서",
        f"{industry} | 목표: {', '.join(objective)}"
    )
    # 기획 개요
    add_text_slide(prs, "1) 제안 개요", 
                   f"- 브랜드: {brand}\n- 업종: {industry}\n- 핵심 목표: {', '.join(objective)}\n\n"
                   "본 제안서는 STP 재정의, 페르소나 구체화, 카피 전략, 채널 믹스 및 KPI 설계를 포함합니다.")
    # STP
    add_text_slide(prs, "2) STP 요약", outputs["stp"])
    # 페르소나
    add_text_slide(prs, "3) 페르소나", outputs["persona"])
    # 카피/슬로건
    add_text_slide(prs, "4) 광고 카피 & 슬로건", outputs["copy"])
    # 매체 전략
    add_text_slide(prs, "5) 매체 전략 & KPI", outputs["media"])
    # 실행 타임라인 vs 예산분배 두 칼럼
    add_two_column_slide(prs, "6) 실행 로드맵",
                         "• 주차별 액션 요약\n- 1주차: 세팅/픽셀\n- 2주차: A/B 테스트\n- 3주차: 스케일/리타겟\n- 4주차: 리포트/리파인",
                         "• 예산 배분(가이드)\n- 상단 퍼널: 30%\n- 중단 퍼널: 40%\n- 하단 퍼널: 30%")
    # 마무리
    add_text_slide(prs, "7) 기대 효과",
                   "• CTR 상승 / CPC 하락\n• CPA 안정화 및 ROAS 개선\n• 크리에이티브-타깃-퍼널 일관성으로 전환 효율 극대화")

    # 간단한 색상 포인트(제목 색)
    for slide in prs.slides:
        try:
            title_shape = slide.shapes.title
            for p in title_shape.text_frame.paragraphs:
                for run in p.runs:
                    run.font.color.rgb = primary
        except Exception:
            pass

    bio = BytesIO()
    prs.save(bio)
    bio.seek(0)
    return bio

# ------------------ 메인 UI ------------------
st.title("📊 마케터용 자동 제안서 도우미")
st.write("업종·타깃·톤·매체를 선택하고 제안을 자동 생성한 뒤, PPT로 바로 내보내자.")

col1, col2 = st.columns([1,1])
with col1:
    st.subheader("🛠️ 생성")
    disabled = not api_key
    generate_btn = st.button("🚀 제안 생성하기", use_container_width=True, disabled=disabled)
with col2:
    st.subheader("📥 내보내기")
    export_btn_disabled = st.session_state.outputs is None
    if export_btn_disabled:
        st.download_button("👉 제안서로 내보내기 (PPTX)", data=b"", file_name="proposal.pptx",
                           disabled=True, use_container_width=True)
    else:
        ppt = build_ppt(brand, industry, objective, st.session_state.outputs)
        st.download_button("👉 제안서로 내보내기 (PPTX)", data=ppt, file_name=f"{brand}_proposal.pptx",
                           mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                           use_container_width=True)

st.divider()

# 생성 실행
if generate_btn:
    client = OpenAI(api_key=api_key)
    kwargs = dict(
        brand=brand, industry=industry, age_range=age_range, gender=gender,
        tone=tone, channels=channels, objective=objective, cta_style=cta_style
    )
    outputs = generate_all(client, **kwargs)
    st.session_state.outputs = outputs
    st.success("생성이 완료되었어. 아래 탭에서 결과를 확인하고, 필요하면 PPT로 내보내!")

# 결과 미리보기
if st.session_state.outputs:
    tabs = st.tabs(["STP", "페르소나", "카피 & 슬로건", "매체 전략/KPI"])
    with tabs[0]:
        st.markdown(st.session_state.outputs["stp"])
    with tabs[1]:
        st.markdown(st.session_state.outputs["persona"])
    with tabs[2]:
        st.markdown(st.session_state.outputs["copy"])
    with tabs[3]:
        st.markdown(st.session_state.outputs["media"])
else:
    st.info("좌측에서 설정을 입력하고 **제안 생성하기**를 눌러 시작하자.", icon="✨")
